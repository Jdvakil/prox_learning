#!/usr/bin/env python3
"""Compose recorded RGB frames and calibrated range returns; no scene rendering."""
from pathlib import Path
import base64, html, json, zipfile
import cairosvg
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT/'images/first_page'
ASSETS = OUT/'recorded_assets'
W, H = 1284, 352
TEAL = '#00A6A6'
PATCH_SENSORS = {'link6_sensor_0', 'link6_sensor_2', 'link6_sensor_3', 'link6_sensor_4'}
PANELS = [
    dict(name='Scene context', source='exo_camera_1.png', crop=[80,0,500,352], x=0),
    dict(name='Wrist RGB', source='wrist_camera.png', crop=[102,0,522,352], x=432),
    dict(name='Measured proximity', source='exo_camera_1.png', crop=[80,0,500,352], x=864),
]
CAPTION = (
    'Complementary views of near-contact manipulation. Synchronized frames from a '
    'recorded fume-hood demonstration show the robot and its surroundings (left), '
    'the wrist-camera view of the grasped cup and nearby clutter (middle), and '
    'local geometry sampled by the robot’s proximity sensors (right; teal points '
    'projected onto the external RGB frame). Fine rays connect measured returns '
    'to their calibrated sensor origins; selected patches connect adjacent pixels '
    'of the sensors’ 8×8 grids. RGB provides visual task context; '
    'proximity supplies direct local range measurements. The paper asks whether '
    'live proximity at inference reduces collisions when task-relevant hazards '
    'are unavailable to RGB. This simulated dataset example illustrates the '
    'sensing inputs, not a policy comparison or evidence that every hazard in '
    'this frame is invisible to every camera.'
)
NOTES = '''EDITING
Import first_page_canva.pptx into Canva. Slide 1 has no text; slide 2 has three editable labels. Each frame is independent and retains its complete source image inside an editable crop. Measured returns and projection rays are separate editable groups. Keep it aligned with its image when resizing. Canva import has not been tested in Canva itself.

SOURCES
All panels use synchronized recorded simulation data from v1011d, row 000_187ba0ce76cb3011, traj_0, frame 150 (9.9 s). This is a scripted dataset demonstration, not a learned-policy rollout. recorded_assets/source_manifest.json records paths, SHA-256 hashes and processing. No scene was newly rendered, no image was generated, and no RGB content was removed, recoloured, retouched or enhanced. Only rectangular crops and layout scaling are applied. RGB sources are 624 x 352.

MEASUREMENTS
Teal markers are recorded proximity samples, pooled by mean over four stored samples. All 2,560 finite samples >= 0.02 m are retained (40 sensors × 8 × 8); 630 lie within the external-camera figure crop. The previous 0.60 m display cutoff is removed; the recorded simulation ranges reach 3.17 m. This displays raw simulated ranges and does not establish physical hardware operating range. Thin rays start at recorded sensor poses. Four wrist patches connect neighboring grid samples only where their range difference is below 7 cm; these connecting lines are visual aids, not extra measurements. Samples outside the figure crop remain in the NPZ and PLY files. Backprojection uses recorded per-sensor intrinsics and transforms. Stored CV extrinsics are verified against cam2world_gl. Camera intrinsics are adjusted for the sidecar MP4 dimensions at the same vertical FOV. No RGB-depth occlusion test is applied, so markers can show geometry behind a visible surface: this is an overlay of sensor coverage, not an RGB segmentation. Uniform marker colour does not encode range. Finite sensor resolution and sample pooling limit reconstruction accuracy. The NPZ and PLY contain all 2,560 pooled ranges and world points; JSON contains projected points before figure cropping.

EXPORTS
Default PNG/SVG/PDF has no text. The labelled variant adds only three panel labels. PDF is 7 inches wide, with vector range markers. Native PNG is 1284 x 352; the 3600-pixel PNG is a layout export and adds no source detail. Source frames, measurements and crops are in recorded_assets. task_start, task_transfer and task_place are frames 0, 250 and 380 from the same external-camera video. The earlier conceptual figure has been replaced.
'''


def returns():
    projected = json.loads((ASSETS/'projected_returns.json').read_text())['exo_camera_1']
    l,t,r,b = PANELS[2]['crop']
    return [dict(x=p['u']-l+864, y=p['v']-t, distance_m=p['distance_m'], sensor=p['sensor'],
                 pixel=p['pixel'], origin_x=p['origin_u']-l+864,
                 origin_y=p['origin_v']-t, origin_z=p['origin_z'])
            for p in sorted(projected, key=lambda a: -a['z'])
            if l+2 <= p['u'] < r-2 and t+2 <= p['v'] < b-2]


def segments(points):
    lines = []
    for p in points:
        if p['origin_z'] > .02 and 865 <= p['origin_x'] <= 1283 and 1 <= p['origin_y'] <= 351:
            lines.append(dict(x1=p['origin_x'], y1=p['origin_y'], x2=p['x'], y2=p['y'],
                              width=.40, opacity=.12, kind='ray', sensor=p['sensor']))
    lookup = {(p['sensor'], *p['pixel']): p for p in points}
    for p in points:
        if p['sensor'] not in PATCH_SENSORS:
            continue
        u,v = p['pixel']
        for key in [(p['sensor'],u+1,v), (p['sensor'],u,v+1)]:
            q = lookup.get(key)
            # Do not bridge depth discontinuities into a fictitious surface.
            if q and abs(p['distance_m']-q['distance_m']) < .07:
                lines.append(dict(x1=p['x'], y1=p['y'], x2=q['x'], y2=q['y'],
                                  width=.55, opacity=.48, kind='grid', sensor=p['sensor']))
    return lines


def svg(labelled, points):
    top = 34 if labelled else 0
    chunks = [f'<svg xmlns="http://www.w3.org/2000/svg" xmlns:xlink="http://www.w3.org/1999/xlink" width="{W}" height="{H+top}" viewBox="0 0 {W} {H+top}">',
              '<title>Recorded scene, wrist RGB, and measured proximity</title>',
              f'<desc>{html.escape(CAPTION)}</desc>',
              f'<rect width="{W}" height="{H+top}" fill="white"/>']
    for i, panel in enumerate(PANELS):
        data = base64.b64encode((ASSETS/f'panel_{i+1}.png').read_bytes()).decode()
        chunks.append(f'<image id="recorded-panel-{i+1}" x="{panel["x"]}" y="{top}" width="420" height="352" xlink:href="data:image/png;base64,{data}"/>')
        if labelled:
            chunks.append(f'<text x="{panel["x"]}" y="22" font-family="Arial, Liberation Sans, sans-serif" font-size="20" fill="#22292E">({chr(97+i)}) {panel["name"]}</text>')
    chunks.append('<g id="sensor-projection-rays-and-grid">')
    for e in segments(points):
        chunks.append(f'<line x1="{e["x1"]:.4f}" y1="{e["y1"]+top:.4f}" x2="{e["x2"]:.4f}" y2="{e["y2"]+top:.4f}" stroke="{TEAL}" stroke-width="{e["width"]}" opacity="{e["opacity"]}"/>')
    chunks.append('</g><g id="recorded-range-returns">')
    for i, p in enumerate(points):
        chunks.append(f'<circle id="return-{i}" cx="{p["x"]:.4f}" cy="{p["y"]+top:.4f}" r="1.55" fill="{TEAL}" stroke="white" stroke-width="0.3"/>')
    chunks += ['</g>', '</svg>']
    return '\n'.join(chunks)


def presentation(points):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(12.84), Inches(3.86)
    unit = 9144
    for labelled in (False, True):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        top = 34 if labelled else 17
        for i, panel in enumerate(PANELS):
            pic = slide.shapes.add_picture(str(ASSETS/panel['source']),
                int(panel['x']*unit), int(top*unit), int(420*unit), int(352*unit))
            l,t,r,b = panel['crop']
            pic.crop_left, pic.crop_right = l/624, (624-r)/624
            pic.crop_top, pic.crop_bottom = t/352, (352-b)/352
            pic.name = f'{panel["name"]} — recorded frame 150 (editable crop)'
            if labelled:
                box = slide.shapes.add_textbox(int(panel['x']*unit), 0, int(420*unit), int(30*unit))
                tf = box.text_frame
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.word_wrap, tf.auto_size = False, MSO_AUTO_SIZE.NONE
                tf._txBody.bodyPr.set('anchorCtr', '0')
                run = tf.paragraphs[0].add_run()
                run.text = f'({chr(97+i)}) {panel["name"]}'
                run.font.name, run.font.size = 'Arial', Pt(14.4)
                run.font.color.rgb = RGBColor.from_string('22292E')
        ray_group = slide.shapes.add_group_shape()
        ray_group.name = 'Calibrated sensor projection rays and 8×8 grid edges'
        ray_group.shapes.turbo_add_enabled = True
        for e in segments(points):
            line = ray_group.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                int(e['x1']*unit), int((e['y1']+top)*unit),
                int(e['x2']*unit), int((e['y2']+top)*unit))
            line.name = f'{e["sensor"]}: {e["kind"]}'
            line.line.color.rgb = RGBColor.from_string(TEAL[1:])
            line.line.width = Pt(e['width']*.72)
            alpha = OxmlElement('a:alpha')
            alpha.set('val', str(round(e['opacity']*100000)))
            line.line._get_or_add_ln().find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr').append(alpha)
        group = slide.shapes.add_group_shape()
        group.name = 'Measured range returns — ungroup to edit points'
        group.shapes.turbo_add_enabled = True
        for p in points:
            shape = group.shapes.add_shape(MSO_SHAPE.OVAL,
                int((p['x']-1.55)*unit), int((p['y']+top-1.55)*unit),
                int(3.1*unit), int(3.1*unit))
            shape.name = f'{p["sensor"]}: {p["distance_m"]:.4f} m'
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor.from_string(TEAL[1:])
            shape.line.color.rgb, shape.line.width = RGBColor(255,255,255), Pt(.216)
            for child in list(shape._element):
                if child.tag.endswith('}style'):
                    shape._element.remove(child)
            shape._element.spPr.append(OxmlElement('a:effectLst'))
        slide.notes_slide.notes_text_frame.text = CAPTION+'\n\n'+NOTES
    prs.core_properties.title = 'Vision and proximity — recorded fume-hood demonstration'
    prs.core_properties.subject = 'Recorded dataset images with editable measured-range overlay'
    prs.save(OUT/'first_page_canva.pptx')


def main():
    for i, panel in enumerate(PANELS):
        Image.open(ASSETS/panel['source']).crop(panel['crop']).save(ASSETS/f'panel_{i+1}.png')
    points = returns()
    for labelled, suffix in [(False,''), (False,'_no_text'), (True,'_labelled')]:
        src = svg(labelled, points)
        (OUT/f'first_page{suffix}.svg').write_text(src)
        cairosvg.svg2pdf(bytestring=src.encode(), write_to=str(OUT/f'first_page{suffix}.pdf'), output_width=672)
        cairosvg.svg2png(bytestring=src.encode(), write_to=str(OUT/f'first_page{suffix}.png'), output_width=3600)
        if not suffix:
            cairosvg.svg2png(bytestring=src.encode(), write_to=str(OUT/'first_page_native.png'))
            panel = src.replace('width="1284" height="352" viewBox="0 0 1284 352"',
                                'width="420" height="352" viewBox="864 0 420 352"', 1)
            (OUT/'pointcloud_panel.svg').write_text(panel)
            cairosvg.svg2png(bytestring=panel.encode(), write_to=str(OUT/'pointcloud_panel.png'), output_width=1800)
            cairosvg.svg2pdf(bytestring=panel.encode(), write_to=str(OUT/'pointcloud_panel.pdf'), output_width=420)
    presentation(points)
    (OUT/'layout.json').write_text(json.dumps(dict(width=W, height=H, panels=PANELS,
        marker_radius=1.55, marker_colour=TEAL, returns=points, segments=segments(points)), indent=2))
    (OUT/'caption.txt').write_text(CAPTION+'\n')
    (OUT/'figure_notes.txt').write_text(NOTES)
    (OUT/'index.html').write_text('<!doctype html><meta charset="utf-8"><title>Recorded paper figure</title>'
        '<style>body{font:16px Arial;max-width:1400px;margin:40px auto;padding:20px;color:#22292e}img{width:100%;margin:24px 0}a{margin-right:24px;color:#007f80}p{line-height:1.6}</style>'
        '<h1>Vision + proximity: recorded experiment imagery</h1><img src="first_page.png" alt="Recorded external camera, wrist camera, and calibrated proximity returns">'
        '<p><a href="first_page_canva.pptx">Editable Canva / PowerPoint</a><a href="first_page.pdf">Paper PDF</a><a href="first_page.png">Text-free PNG</a><a href="first_page.svg">SVG</a><a href="first_page_labelled.png">Minimal labels</a><a href="canva_assets.zip">Asset bundle</a></p>'
        '<p>'+html.escape(CAPTION)+'</p><img src="first_page_labelled.png" alt="Variant with three panel labels">')
    with zipfile.ZipFile(OUT/'canva_assets.zip','w',zipfile.ZIP_DEFLATED) as z:
        for p in [OUT/'first_page_canva.pptx', OUT/'first_page.svg', OUT/'pointcloud_panel.svg', OUT/'pointcloud_panel.png', OUT/'caption.txt', OUT/'figure_notes.txt', *ASSETS.glob('*')]:
            z.write(p,p.relative_to(OUT))
    print(f'Built recording figures with {len(points)} editable measured markers.')


if __name__ == '__main__':
    main()
